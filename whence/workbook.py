# SPDX-License-Identifier: Apache-2.0
"""Spreadsheets: the last mile, where most figures actually get made.

A close does not end in a warehouse. It ends in a workbook somebody built by
pasting exports into tabs and writing formulas over them, and the number on the
board slide is whatever that workbook says. This module reads such a workbook
from the .xlsx file alone -- no Excel, no add-in, nothing running on the
analyst's machine -- and records the figure it produces with the same kind of
provenance the Python close produces.

It does that in two layers, and the distinction is worth keeping straight:

  1. **The pasted inputs.** Tabs that hold the close's source tables (orders,
     pricebook, FX, adjustments) are read as tables, cell by cell, into the same
     text the CSV reader produces -- so they mint the *same source facts*, with
     the same ids. Then the reference close (pipeline.close, unchanged) is run
     over them. If the workbook's inputs are the CSVs' inputs, this lands on the
     CSV close's id exactly: identity survives the trip into a spreadsheet.

  2. **The analyst's formulas.** Starting from the workbook's headline cell,
     every formula it depends on is parsed, evaluated by a small evaluator
     (whence/formula.py says why it is small), and lowered into ordinary
     tracked transforms -- `multiply`, `sum`, `subtract` -- over the facts of
     the cells it reads. The result is a figure whose lineage is the
     workbook's own logic, verifiable by the standard verifier with no
     knowledge that it ever came from Excel.

Comparing the two is the point. When they agree, the spreadsheet did what the
close does. When they do not, the difference has a cause, and this module goes
looking for it:

  * **A typed-over formula.** A constant sitting in a column of identical,
    copied-down formulas is almost always a formula someone overwrote. It is
    recorded as `asserted` -- a person typed it -- and its effect on the
    headline is measured by re-evaluating the workbook with the formula put
    back.
  * **A stale or edited cached value.** Excel stores each formula's last
    result beside it. Every formula is recomputed and compared; a mismatch is
    recorded as found, and the figure is minted with the value the workbook
    *claims*, so the standard verifier rejects it too.
  * **A formula outside the evaluator's subset.** Recorded as an opaque node:
    its lineage (which cells it read) is kept, but it is never reported as
    verified.

What is deliberately *not* hashed, by the same rule as SPEC.md section 3.8:
the sheet name, the cell address, and the workbook's digest are where a value
was *seen*, and go in observations. A figure's identity does not change because
someone inserted a row above it.

Values out of cells, precisely (SPEC.md section 9 says the same at length):

  * A number is fixed-point text at the scale its **number format** declares.
    `#,##0.00` is two places, `0.0000` is four, `0` is none. A cell's format is
    how a spreadsheet says what precision a value has, and it is what the
    analyst saw. A float is read through its shortest round-trip decimal form,
    never formatted at a guessed precision.
  * A date is ISO 8601 (`2026-09-30`). openpyxl converts serial dates using
    the workbook's own epoch, so a 1904-system workbook from an old Mac and
    the 1900 system's phantom 29 February 1900 are both handled before a value
    reaches here.
  * A blank cell in a table is the empty string. A spreadsheet cannot tell a
    blank from an empty string, so neither can this.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from . import arith
from .core import (CanonicalizationError, Fact, MemoryStore, Store,
                   asserted_fact, quantize_text, source_fact)
from .formula import Ref, Unsupported, a1, col_letter, normalize, parse
from .pipeline import RunResult, close, line_revenue

__all__ = ["number_scale", "cell_text", "Book", "SheetTable", "Finding",
           "WorkbookResult", "ingest_workbook", "SUPPORTED_FUNCTIONS"]

#: The functions the evaluator understands. Anything else is recorded as an
#: opaque node -- kept in the lineage, never reported as verified.
SUPPORTED_FUNCTIONS = ("SUM", "SUMIFS", "VLOOKUP", "ROUND")

#: Opaque formulas are recorded under this transform name. No build registers
#: it, so every verifier reports such a figure as not verifiable (section 4).
OPAQUE = "xl.opaque"

_LOWERED = {"multiply": line_revenue, "sum": arith.add,
            "subtract": arith.subtract}

CENT = Decimal("0.01")


# ---------------------------------------------------------------------------
# Cells as text.
# ---------------------------------------------------------------------------

_FMT_NOISE = re.compile(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.')


def number_scale(fmt: Optional[str]) -> Optional[int]:
    """Decimal places a number format displays, or None for General / text.

    Only the first section counts (the positive-number format). Literal text,
    colours and conditions in square brackets, escapes and padding are
    stripped first. A percent format displays the value times 100, so its
    stored value carries two more places than it shows.
    """
    if not fmt or fmt.strip().lower() == "general":
        return None
    section = _FMT_NOISE.sub("", fmt.split(";")[0])
    if not re.search(r"[0#?]", section):
        return None
    if "E+" in section.upper() or "E-" in section.upper():
        return None                              # scientific: no fixed scale
    places = 0
    if "." in section:
        places = len(re.match(r"[0#?]*", section.split(".", 1)[1]).group(0))
    return places + (2 if "%" in section else 0)


def _decimal(value: Any) -> Decimal:
    """A cell number as an exact Decimal. A float goes through repr, which is
    the shortest text that round-trips to the same double -- so 0.1 is 0.1,
    not 0.1000000000000000055511151231257827."""
    if isinstance(value, bool):
        raise CanonicalizationError("a boolean is not a number")
    if isinstance(value, float):
        return Decimal(repr(value))
    return Decimal(value)


def number_text(value: Any, fmt: Optional[str]) -> str:
    """Canonical text for a number in a cell with this format."""
    d = _decimal(value)
    scale = number_scale(fmt)
    if scale is None:
        if d == d.to_integral_value():
            return str(int(d))
        return format(d.normalize(), "f")
    return quantize_text(d, scale)


def cell_text(value: Any, fmt: Optional[str] = None) -> str:
    """One spreadsheet cell as the text a CSV export of it would contain."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, (int, float, Decimal)):
        return number_text(value, fmt)
    if isinstance(value, str):
        return value
    raise CanonicalizationError(
        "no canonical text for a %s cell" % type(value).__name__)


# ---------------------------------------------------------------------------
# The workbook.
# ---------------------------------------------------------------------------

class Book:
    """An .xlsx file, opened twice: once for formulas, once for the values
    Excel cached beside them."""

    def __init__(self, path: Path | str):
        import openpyxl                                  # optional dependency
        self.path = Path(path)
        raw = self.path.read_bytes()
        self.digest = hashlib.sha256(raw).hexdigest()
        self.name = self.path.name
        self.formulas = openpyxl.load_workbook(self.path, data_only=False)
        self.values = openpyxl.load_workbook(self.path, data_only=True)
        props = self.formulas.properties
        self.author = props.lastModifiedBy or props.creator or "unknown"
        self.modified = (props.modified.date().isoformat()
                         if props.modified else "")

    def sheet(self, name: str):
        for ws in self.formulas.worksheets:
            if ws.title.lower() == name.lower():
                return ws
        raise Unsupported("a reference to a sheet that does not exist (%s)"
                          % name)

    def cached(self, sheet: str, row: int, col: int) -> Any:
        return self.values[self.sheet(sheet).title].cell(row, col).value

    def headline(self) -> Tuple[str, int, int]:
        """The cell named `Headline`, if the workbook defines one."""
        names = self.formulas.defined_names
        for key in list(names.keys()):
            if key.lower() == "headline":
                for sheet, coord in names[key].destinations:
                    return _locate(self, "%s!%s" % (sheet, coord))
        raise ValueError(
            "%s defines no cell named Headline; say which cell to follow, "
            "e.g. --cell 'Summary!B10'" % self.name)


def _locate(book: Book, ref: str) -> Tuple[str, int, int]:
    from .formula import parse_ref
    try:
        parsed = parse_ref(ref.replace("$", ""))
    except Unsupported:
        raise ValueError("not a cell reference: %r" % ref) from None
    if parsed.sheet is None or not parsed.is_cell:
        raise ValueError("give a single cell with its sheet, e.g. "
                         "'Summary!B10' (got %r)" % ref)
    return book.sheet(parsed.sheet).title, parsed.r1, parsed.c1


# ---------------------------------------------------------------------------
# Layer 1: pasted tables, shaped like pipeline.SourceFile.
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class SheetTable:
    """A tab holding a table, shaped so pipeline.close cannot tell it from a
    CSV. Header in row 1; one record per row until the first empty row."""

    name: str                       # "q3_close.xlsx [Pricebook]"
    sheet: str
    digest: str                     # of the workbook file
    rows: List[dict]
    columns: Dict[str, int]

    def where(self, **match) -> dict:
        for row in self.rows:
            if all(row.get(k) == v for k, v in match.items()):
                return row
        raise KeyError("no row in %s matching %r" % (self.name, match))

    def coordinates(self, row, column: str) -> dict:
        return {"file": self.name, "row": row["_line"], "column": column,
                "cell": "%s!%s%s" % (self.sheet,
                                      col_letter(self.columns[column]),
                                      row["_line"]),
                "digest": self.digest}


def read_table(book: Book, sheet: str) -> SheetTable:
    ws = book.sheet(sheet)
    cached = book.values[ws.title]
    header = {}
    for col in range(1, ws.max_column + 1):
        name = ws.cell(1, col).value
        if isinstance(name, str) and name.strip():
            header[name.strip()] = col
    rows = []
    for r in range(2, ws.max_row + 1):
        cells = {name: (ws.cell(r, c), cached.cell(r, c).value)
                 for name, c in header.items()}
        if all(v is None for _cell, v in cells.values()):
            break
        row = {name: cell_text(v, cell.number_format)
               for name, (cell, v) in cells.items()}
        row["_line"] = str(r)
        rows.append(row)
    return SheetTable(name="%s [%s]" % (book.name, ws.title), sheet=ws.title,
                      digest=book.digest, rows=rows, columns=header)


#: The tabs the reference close needs, and the columns each must have.
INPUT_TABS = {
    "orders": ("Orders", {"order_id", "order_date", "sku", "qty", "region",
                          "currency", "unit_price_at_sale"}),
    "pricebook": ("Pricebook", {"sku", "list_price", "effective_date",
                                "last_modified_by", "approval_ref"}),
    "fx": ("FX", {"rate_date", "currency", "usd_per_unit"}),
    "adjustments": ("Adjustments", {"adj_id", "entry_date", "amount_usd",
                                    "reason", "author", "source"}),
}


def input_tables(book: Book) -> Optional[Dict[str, SheetTable]]:
    """The four input tabs, if this workbook has all of them; else None."""
    found = {}
    for role, (sheet, needed) in INPUT_TABS.items():
        try:
            table = read_table(book, sheet)
        except Unsupported:
            return None
        if not needed <= set(table.columns):
            return None
        found[role] = table
    return found


# ---------------------------------------------------------------------------
# Layer 2: the formulas.
# ---------------------------------------------------------------------------

@dataclass
class V:
    """An evaluated value, and where it came from.

    Exactly one of three things backs a number: a `fact` (the value *is* that
    fact's value), a `pending` operation not yet minted -- (transform,
    operands, rounded) -- or nothing, for a literal written into the formula.
    """
    value: Any
    fact: Optional[Fact] = None
    pending: Optional[tuple] = None
    literal: bool = False

    @property
    def numeric(self) -> bool:
        return isinstance(self.value, Decimal) and not isinstance(self.value,
                                                                  bool)


@dataclass
class RangeV:
    """A rectangle of cells, evaluated only as far as something needs."""
    ev: "Evaluator"
    sheet: str
    ref: Ref

    def cells(self):
        for r in range(self.ref.r1, self.ref.r2 + 1):
            for c in range(self.ref.c1, self.ref.c2 + 1):
                yield r, c

    def get(self, r: int, c: int) -> V:
        return self.ev.cell(self.sheet, r, c)

    def peek(self, r: int, c: int) -> Any:
        return self.ev.peek(self.sheet, r, c)


@dataclass
class Finding:
    """Something a reviewer should know about the workbook."""
    kind: str                       # override | cache | opaque | uncached
    cell: str                       # "Revenue!D57"
    message: str
    detail: Dict[str, str] = field(default_factory=dict)


class Evaluator:
    """Evaluates cells on demand, minting a Fact for every number it touches.

    Lazy on purpose: only the headline's dependencies are visited, so a
    workbook's lineage is the lineage of the figure it reports, not of every
    scratch calculation someone left in column Z.
    """

    def __init__(self, book: Book, store: Store,
                 inputs: Dict[Tuple[str, str], Fact],
                 restore_overrides: bool = False):
        self.book = book
        self.store = store
        self.inputs = inputs
        self.restore = restore_overrides
        self.memo: Dict[Tuple[str, int, int], V] = {}
        self.active: set = set()
        self.findings: List[Finding] = []
        self.formula_cells = 0

    # -- naming --------------------------------------------------------------

    def where(self, sheet: str, row: int, col: int) -> str:
        return "%s!%s" % (sheet, a1(row, col))

    def _header(self, sheet: str, col: int) -> str:
        v = self.peek(sheet, 1, col) if col else None
        return v.strip() if isinstance(v, str) else ""

    def identity(self, sheet: str, row: int, col: int) -> Tuple[str, str]:
        """(key, label) for a cell, in the sheet's own terms.

        The row is named by the text in column A of that row -- an order id,
        a line label -- and the column by its header in row 1. Neither is a
        cell address, so the key survives rows being inserted above. A cell
        with neither falls back to its address, which is all there is.
        """
        rowname = ""
        if col != 1 and row != 1:
            v = self.peek(sheet, row, 1)
            if isinstance(v, str) and v.strip():
                rowname = v.strip()
        header = self._header(sheet, col) if row != 1 else ""
        if not rowname:
            return ("xl:%s" % self.where(sheet, row, col),
                    ("%s %s" % (self.where(sheet, row, col), header)).strip())
        key = "xl:%s:%s" % (sheet, rowname) + (".%s" % header if header else "")
        return key, ("%s %s" % (rowname, header)).strip()

    # -- cells ---------------------------------------------------------------

    def peek(self, sheet: str, row: int, col: int) -> Any:
        """A cell's plain value, for comparing -- lookup keys, criteria --
        without minting anything for it. A formula is evaluated."""
        ws = self.book.sheet(sheet)
        raw = ws.cell(row, col).value
        if isinstance(raw, str) and raw.startswith("="):
            return self.cell(ws.title, row, col).value
        if isinstance(raw, (int, float)) and not isinstance(raw, bool):
            return _decimal(raw)
        return raw

    def cell(self, sheet: str, row: int, col: int) -> V:
        sheet = self.book.sheet(sheet).title
        key = (sheet, row, col)
        if key in self.memo:
            return self.memo[key]
        if key in self.active:
            raise Unsupported("a circular reference through %s"
                              % self.where(*key))
        self.active.add(key)
        try:
            raw = self.book.sheet(sheet).cell(row, col).value
            if isinstance(raw, str) and raw.startswith("="):
                value = self._formula_cell(sheet, row, col, raw)
            elif raw is not None and not isinstance(raw, (str, int, float,
                                                          bool, date, time)):
                # openpyxl hands back objects for array formulas and the like.
                value = self._opaque(sheet, row, col, str(raw),
                                     "a formula type this build cannot read")
            else:
                value = self._constant_cell(sheet, row, col, raw)
        finally:
            self.active.discard(key)
        self.memo[key] = value
        return value

    def _constant_cell(self, sheet: str, row: int, col: int, raw: Any) -> V:
        if not isinstance(raw, (int, float)) or isinstance(raw, bool):
            return V(raw)                       # text, a date, blank: no fact
        fmt = self.book.sheet(sheet).cell(row, col).number_format
        text = number_text(raw, fmt)

        known = self.inputs.get((sheet, a1(row, col)))
        if known is not None:
            return V(Decimal(known.value), fact=known)

        pattern = self._column_formula(sheet, row, col)
        if pattern is not None:
            return self._typed_over(sheet, row, col, text, pattern)

        key, label = self.identity(sheet, row, col)
        fact = source_fact(text, {"key": "src:" + key, "label": label})
        self.store.put(fact)
        self._observe(fact, sheet, row, col)
        return V(Decimal(text), fact=fact)

    def _column_formula(self, sheet: str, row: int, col: int):
        """If this constant sits in a run of identical, copied-down formulas,
        return (neighbour row, neighbour formula). Else None.

        Both nearest non-blank neighbours in the column must be formulas and
        must normalize to the same text; at the edge of the data, the one
        neighbour there is must, and the one beyond it must match it too.
        That is a strict test on purpose -- a constant in a column of
        constants, or between two different formulas, is just a value.
        """
        ws = self.book.sheet(sheet)

        def neighbour(step):
            r = row + step
            while 2 <= r <= ws.max_row:
                v = ws.cell(r, col).value
                if v is not None:
                    return r, v
                r += step
            return None

        def norm(found):
            r, v = found
            if isinstance(v, str) and v.startswith("="):
                return normalize(v, r, col)
            return None

        up, down = neighbour(-1), neighbour(1)
        if up is None and down is None:
            return None
        sides = [s for s in (up, down) if s is not None]
        forms = [norm(s) for s in sides]
        if any(f is None for f in forms) or len(set(forms)) != 1:
            return None
        if len(sides) == 1:
            # At the edge of the data: one neighbour is not a pattern, two
            # matching ones are.
            r2 = sides[0][0] + (-1 if up else 1)
            v2 = ws.cell(r2, col).value if 2 <= r2 <= ws.max_row else None
            if not (isinstance(v2, str) and v2.startswith("=")
                    and normalize(v2, r2, col) == forms[0]):
                return None
        return sides[0]

    def _typed_over(self, sheet, row, col, text, pattern) -> V:
        from openpyxl.formula.translate import Translator
        src_row, src_formula = pattern
        formula = Translator(src_formula, origin=a1(src_row, col)) \
            .translate_formula(a1(row, col))
        coord = self.where(sheet, row, col)

        if self.restore:
            return self._formula_cell(sheet, row, col, formula, cached=False)

        try:
            would = self.evaluate(formula, sheet, row, col).value
        except Unsupported:
            would = None
        key, label = self.identity(sheet, row, col)
        fact = asserted_fact(text, {
            "key": key,
            "label": "%s (typed over its formula)" % label,
            "actor": self.book.author,
            "reason": "%s holds a typed number where its column has a formula"
                      % coord,
            "evidence": "workbook last saved by %s" % self.book.author,
            "xl.formula": normalize(formula, row, col),
        })
        self.store.put(fact)
        self._observe(fact, sheet, row, col)
        detail = {"typed": text, "formula": formula,
                  "author": self.book.author}
        if isinstance(would, Decimal):
            detail["formula gives"] = quantize_text(would)
        self.findings.append(Finding(
            kind="override", cell=coord,
            message="%s is a typed number in a column of formulas" % coord,
            detail=detail))
        return V(Decimal(text), fact=fact)

    def _formula_cell(self, sheet, row, col, formula, cached=True) -> V:
        self.formula_cells += 1
        coord = self.where(sheet, row, col)
        norm = normalize(formula, row, col)
        try:
            v = self.evaluate(formula, sheet, row, col)
            if isinstance(v, RangeV):
                raise Unsupported("a formula whose result is a range")
            if v.pending is not None:
                self._check_exact(v)
        except Unsupported as exc:
            return self._opaque(sheet, row, col, formula, str(exc))

        if not v.numeric:
            return V(v.value)                   # text out of a formula: no fact

        key, label = self.identity(sheet, row, col)
        meta = {"key": key, "label": label, "xl.formula": norm}

        if v.pending is not None:
            fact = self._mint(v, meta, top=True)
        elif v.fact is not None:
            fact = v.fact                       # =B4, or a lookup: same fact
        else:
            fact = self._literal(v, meta, coord, 0)

        # With overrides restored, the saved values downstream of them are
        # expected to disagree -- they were computed from the typed number --
        # so the counterfactual run does not compare against them.
        if cached and not self.restore:
            fact = self._compare_cache(fact, sheet, row, col, minted=v.fact is None)
        if v.fact is None:
            self._observe(fact, sheet, row, col)
        return V(Decimal(fact.value), fact=fact)

    def _compare_cache(self, fact, sheet, row, col, minted: bool) -> Fact:
        """Compare the evaluator's answer with what Excel saved for the cell.

        For a figure this module minted, a disagreement is recorded with the
        value the *workbook* claims, so the fact says what the spreadsheet
        says -- and replaying its transform, which the standard verifier does,
        produces a different number. The tampering is then caught by the same
        check that catches it everywhere else, not by a special case here.
        """
        coord = self.where(sheet, row, col)
        raw = self.book.cached(sheet, row, col)
        if raw is None:
            self.findings.append(Finding(
                kind="uncached", cell=coord,
                message="%s has no saved value; Whence computed it" % coord))
            return fact
        if not isinstance(raw, (int, float)) or isinstance(raw, bool):
            return fact
        scale = len(fact.value.partition(".")[2])
        claimed = quantize_text(_decimal(raw), scale)
        if claimed == fact.value:
            return fact
        self.findings.append(Finding(
            kind="cache", cell=coord,
            message="%s shows %s but its formula gives %s"
                    % (coord, claimed, fact.value),
            detail={"saved": claimed, "recomputed": fact.value}))
        if not minted:
            return fact
        claimed_fact = Fact(value=claimed, kind=fact.kind,
                            transform=fact.transform, inputs=fact.inputs,
                            metadata=dict(fact.metadata))
        self.store.put(claimed_fact)
        return claimed_fact

    def _opaque(self, sheet, row, col, formula, why) -> V:
        """A formula outside the subset. Kept, labelled, never verified."""
        coord = self.where(sheet, row, col)
        raw = self.book.cached(sheet, row, col)
        if not isinstance(raw, (int, float)) or isinstance(raw, bool):
            # A text or blank result has nothing to record as a figure.
            self.findings.append(Finding(
                kind="opaque", cell=coord,
                message="%s could not be evaluated: %s" % (coord, why),
                detail={"formula": formula}))
            return V(raw)

        fmt = self.book.sheet(sheet).cell(row, col).number_format
        text = number_text(raw, fmt)
        parents = []
        try:
            for ref in _refs(parse(formula)):
                parents.extend(self._ref_facts(sheet, ref))
        except Unsupported:
            pass
        parents = list(dict.fromkeys(f.id for f in parents))

        key, label = self.identity(sheet, row, col)
        meta = {"key": key, "label": label, "xl.reason": why}
        try:
            meta["xl.formula"] = normalize(formula, row, col)
        except Unsupported:
            meta["xl.formula"] = formula
        if parents:
            fact = Fact(value=text, kind="derived", transform=OPAQUE,
                        inputs=tuple(parents), metadata=meta)
        else:
            meta.update({"actor": self.book.author,
                         "reason": "formula in %s could not be evaluated: %s"
                                   % (coord, why)})
            fact = asserted_fact(text, meta)
        self.store.put(fact)
        self._observe(fact, sheet, row, col)
        self.findings.append(Finding(
            kind="opaque", cell=coord,
            message="%s could not be evaluated: %s" % (coord, why),
            detail={"formula": formula}))
        return V(Decimal(text), fact=fact)

    def _ref_facts(self, sheet, ref: Ref) -> List[Fact]:
        out = []
        target = ref.sheet or sheet
        for r in range(ref.r1, ref.r2 + 1):
            for c in range(ref.c1, ref.c2 + 1):
                try:
                    v = self.cell(target, r, c)
                except Unsupported:
                    continue
                if v.fact is not None:
                    out.append(v.fact)
        return out

    def _observe(self, fact: Fact, sheet: str, row: int, col: int) -> None:
        self.store.observe(fact.id, file="%s [%s]" % (self.book.name, sheet),
                           row=str(row), column=self._header(sheet, col),
                           cell=self.where(sheet, row, col),
                           digest=self.book.digest)

    # -- minting -------------------------------------------------------------

    def _check_exact(self, v: V) -> None:
        """Money is recorded at a declared scale (SPEC.md section 3.7). A
        result that is not exact to the cent, and was not rounded by the
        formula itself, has no honest two-place form -- so rather than round
        it silently, the cell is treated as opaque."""
        transform, operands, rounded = v.pending
        for op in operands:
            if op.pending is not None:
                self._check_exact(op)
        if not rounded and v.value != v.value.quantize(CENT):
            raise Unsupported(
                "a result that is not exact to the cent and is not rounded "
                "(%s)" % v.value)

    def _mint(self, v: V, meta: dict, top: bool = False) -> Fact:
        transform, operands, _rounded = v.pending
        parents = []
        for i, op in enumerate(operands, start=1):
            if op.fact is not None:
                parents.append(op.fact)
            elif op.pending is not None:
                sub = dict(meta, key="%s#%d" % (meta["key"], i),
                           label="%s, step %d" % (meta["label"], i))
                parents.append(self._mint(op, sub))
            else:
                parents.append(self._literal(op, meta, meta["key"], i))
        fn = _LOWERED[transform]
        if transform == "sum":
            return fn(parents, meta=meta, store=self.store)
        return fn(*parents, meta=meta, store=self.store)

    def _literal(self, v: V, meta: dict, where: str, i: int) -> Fact:
        """A number typed into a formula -- `=B4*1.08` -- is an assumption a
        person made. It is recorded as asserted, like any other typed value."""
        text = format(v.value.normalize(), "f") if v.value != 0 else "0"
        fact = asserted_fact(text, {
            "key": "%s.const%d" % (meta["key"], i),
            "label": "%s written into a formula" % text,
            "actor": self.book.author,
            "reason": "constant typed into the formula behind %s"
                      % meta["label"],
        })
        self.store.put(fact)
        return fact

    # -- the evaluator ---------------------------------------------------------

    def evaluate(self, formula: str, sheet: str, row: int, col: int):
        return self._eval(parse(formula), sheet)

    def _eval(self, node, sheet):
        kind = node[0]
        if kind == "num":
            return V(node[1], literal=True)
        if kind == "str":
            return V(node[1])
        if kind == "bool":
            return V(node[1])
        if kind == "ref":
            ref = node[1]
            target = self.book.sheet(ref.sheet).title if ref.sheet else sheet
            if ref.is_cell:
                return self.cell(target, ref.r1, ref.c1)
            return RangeV(self, target, ref)
        if kind == "neg":
            x = self._scalar(self._eval(node[1], sheet))
            if x.literal:
                return V(-x.value, literal=True)
            raise Unsupported("negating a cell")
        if kind == "bin":
            return self._binary(node[1], self._scalar(self._eval(node[2], sheet)),
                                self._scalar(self._eval(node[3], sheet)))
        if kind == "fn":
            name, args = node[1], node[2]
            handler = getattr(self, "_fn_" + name.lower(), None)
            if name not in SUPPORTED_FUNCTIONS or handler is None:
                raise Unsupported("%s(), which is outside the supported "
                                  "functions (%s)"
                                  % (name, ", ".join(SUPPORTED_FUNCTIONS)))
            return handler([self._eval(a, sheet) for a in args])
        raise Unsupported(kind)

    def _scalar(self, x):
        if isinstance(x, RangeV):
            raise Unsupported("a range used as a single value")
        return x

    def _number(self, x: V, what: str) -> V:
        if not x.numeric:
            raise Unsupported("%s that is %s, not a number"
                              % (what, "blank" if x.value is None else
                                 "text" if isinstance(x.value, str) else
                                 type(x.value).__name__))
        return x

    def _binary(self, op: str, a: V, b: V) -> V:
        a, b = self._number(a, "an operand"), self._number(b, "an operand")
        if a.literal and b.literal:
            raise Unsupported("arithmetic on two constants")
        if op == "+":
            ops = []
            for x in (a, b):
                if x.pending and x.pending[0] == "sum" and not x.pending[2]:
                    ops.extend(x.pending[1])
                else:
                    ops.append(x)
            return V(a.value + b.value, pending=("sum", ops, False))
        if op == "-":
            return V(a.value - b.value, pending=("subtract", [a, b], False))
        if op == "*":
            return V(a.value * b.value, pending=("multiply", [a, b], False))
        raise Unsupported("the operator %s" % op)

    @staticmethod
    def _combine(numbers: List[V], what: str) -> V:
        if not numbers:
            raise Unsupported("%s over no numbers" % what)
        if len(numbers) == 1:
            return numbers[0]              # a sum of one thing is that thing
        return V(sum((n.value for n in numbers), Decimal("0")),
                 pending=("sum", list(numbers), False))

    def _fn_sum(self, args) -> V:
        found = []
        for arg in args:
            if isinstance(arg, RangeV):
                for r, c in arg.cells():
                    v = arg.get(r, c)
                    if v.numeric:          # SUM skips text and blanks, as Excel
                        found.append(v)
            else:
                found.append(self._number(arg, "a SUM argument"))
        return self._combine(found, "SUM")

    def _fn_sumifs(self, args) -> V:
        if len(args) < 3 or len(args) % 2 == 0:
            raise Unsupported("SUMIFS with the wrong number of arguments")
        target = args[0]
        pairs = [(args[i], args[i + 1]) for i in range(1, len(args), 2)]
        if not isinstance(target, RangeV) or any(
                not isinstance(rng, RangeV) or rng.ref.shape != target.ref.shape
                for rng, _ in pairs):
            raise Unsupported("SUMIFS over ranges of different shapes")
        wanted = [(rng, _criterion(self._scalar(crit))) for rng, crit in pairs]
        found = []
        for i, (r, c) in enumerate(target.cells()):
            ok = True
            for rng, want in wanted:
                rr, cc = _nth(rng.ref, i)
                if not _matches(rng.peek(rr, cc), want):
                    ok = False
                    break
            if ok:
                v = target.get(r, c)
                if v.numeric:
                    found.append(v)
        return self._combine(found, "SUMIFS")

    def _fn_vlookup(self, args) -> V:
        if len(args) not in (3, 4):
            raise Unsupported("VLOOKUP with the wrong number of arguments")
        key, table, index = self._scalar(args[0]), args[1], self._scalar(args[2])
        if not isinstance(table, RangeV):
            raise Unsupported("VLOOKUP into something that is not a range")
        if len(args) == 4:
            exact = self._scalar(args[3]).value
            if exact not in (False, Decimal(0)):
                raise Unsupported("an approximate-match VLOOKUP")
        else:
            raise Unsupported("an approximate-match VLOOKUP")
        if not (index.literal and index.value == int(index.value)):
            raise Unsupported("a VLOOKUP column that is not a fixed number")
        col = table.ref.c1 + int(index.value) - 1
        if not table.ref.c1 <= col <= table.ref.c2:
            raise Unsupported("a VLOOKUP column outside its table")
        want = _criterion(key)
        for r in range(table.ref.r1, table.ref.r2 + 1):
            if _matches(table.peek(r, table.ref.c1), want):
                return table.get(r, col)
        raise Unsupported("a VLOOKUP that finds nothing (#N/A)")

    def _fn_round(self, args) -> V:
        if len(args) != 2:
            raise Unsupported("ROUND with the wrong number of arguments")
        x = self._number(self._scalar(args[0]), "a ROUND argument")
        places = self._scalar(args[1])
        if not (places.literal and places.value == 2):
            raise Unsupported("ROUND to other than two places")
        # Excel's ROUND rounds halves away from zero, which is ROUND_HALF_UP.
        rounded = x.value.quantize(CENT, rounding=ROUND_HALF_UP)
        if x.pending is not None:
            t, ops, _ = x.pending
            return V(rounded, pending=(t, ops, True))
        if x.literal:
            return V(rounded, literal=True)
        if rounded != x.value:
            raise Unsupported("ROUND that changes a stored value")
        return x


def _nth(ref: Ref, i: int) -> Tuple[int, int]:
    width = ref.c2 - ref.c1 + 1
    return ref.r1 + i // width, ref.c1 + i % width


def _criterion(v: V):
    """What a lookup or SUMIFS criterion matches against. Only equality is
    supported; operators (">100") and wildcards are not."""
    value = v.value
    if isinstance(value, str):
        if value[:1] in "<>=" or "*" in value or "?" in value:
            raise Unsupported("a criterion with an operator or wildcard (%s)"
                              % value)
        return ("text", value.casefold())
    if isinstance(value, datetime):
        return ("date", value.date()) if value.time() == time(0) \
            else ("datetime", value)
    if isinstance(value, date):
        return ("date", value)
    if isinstance(value, bool):
        return ("bool", value)
    if isinstance(value, Decimal):
        return ("number", value)
    raise Unsupported("a blank criterion")


def _matches(value: Any, want) -> bool:
    kind, target = want
    if isinstance(value, bool):
        return kind == "bool" and value == target
    if kind == "text":
        return isinstance(value, str) and value.casefold() == target
    if kind == "date":
        if isinstance(value, datetime):
            return value.time() == time(0) and value.date() == target
        return isinstance(value, date) and value == target
    if kind == "number":
        try:
            return value is not None and _decimal(value) == target
        except (CanonicalizationError, InvalidOperation, TypeError,
                ValueError):
            return False
    return value == target


def _refs(node) -> List[Ref]:
    if node[0] == "ref":
        return [node[1]]
    out = []
    for part in node[1:]:
        if isinstance(part, tuple):
            out.extend(_refs(part))
        elif isinstance(part, list):
            for item in part:
                out.extend(_refs(item))
    return out


# ---------------------------------------------------------------------------
# Ingest.
# ---------------------------------------------------------------------------

@dataclass
class WorkbookResult:
    book: str
    digest: str
    cell: str
    headline: Fact
    findings: List[Finding]
    formula_cells: int
    reference: Optional[RunResult] = None
    override_effect: Optional[Decimal] = None
    restored_value: Optional[str] = None
    ran_at: str = ""

    @property
    def overrides(self) -> List[Finding]:
        return [f for f in self.findings if f.kind == "override"]

    @property
    def cache_mismatches(self) -> List[Finding]:
        return [f for f in self.findings if f.kind == "cache"]

    @property
    def opaque(self) -> List[Finding]:
        return [f for f in self.findings if f.kind == "opaque"]

    @property
    def reference_delta(self) -> Optional[Decimal]:
        if self.reference is None:
            return None
        return Decimal(self.headline.value) - Decimal(self.reference.net.value)


def ingest_workbook(store: Store, path: Path | str, cell: Optional[str] = None,
                    label: Optional[str] = None) -> WorkbookResult:
    """Record the figure a workbook reports, with its full lineage.

    `cell` is the figure to follow ('Summary!B10'); by default the cell the
    workbook names `Headline`.
    """
    book = Book(path)
    sheet, row, col = _locate(book, cell) if cell else book.headline()

    # Layer 1: the pasted inputs, through the reference close.
    reference = None
    known: Dict[Tuple[str, str], Fact] = {}
    tables = input_tables(book)
    if tables is not None:
        reference = close(store, label="reference close over %s" % book.name,
                          **tables)
        for fid, fact in store.closure(reference.net.id).items():
            for obs in store.observations(fid):
                if obs.get("digest") == book.digest and obs.get("cell"):
                    tab, _, coord = obs["cell"].partition("!")
                    known[(tab, coord)] = fact

    # Layer 2: the workbook's own formulas, from the headline down.
    ev = Evaluator(book, store, known)
    head = ev.cell(sheet, row, col)
    if head.fact is None:
        raise ValueError("%s does not hold a number" % ev.where(sheet, row, col))

    result = WorkbookResult(
        book=book.name, digest=book.digest, cell=ev.where(sheet, row, col),
        headline=head.fact, findings=ev.findings,
        formula_cells=ev.formula_cells, reference=reference)

    # What would the headline be with every typed-over formula put back?
    if result.overrides:
        scratch = MemoryStore()
        alt = Evaluator(book, scratch, known, restore_overrides=True)
        restored = alt.cell(sheet, row, col)
        if restored.fact is not None:
            result.restored_value = restored.fact.value
            result.override_effect = (Decimal(head.fact.value)
                                      - Decimal(restored.fact.value))

    result.ran_at = datetime.now(timezone.utc).replace(
        microsecond=0).isoformat()
    store.record_run({
        "ran_at": result.ran_at,
        "label": label or "workbook %s" % book.name,
        "kind": "workbook",
        "root": head.fact.id,
        "cell": result.cell,
        "sources": {book.name: book.digest},
        "reference_root": reference.net.id if reference else "",
        "findings": [{"kind": f.kind, "cell": f.cell, "message": f.message}
                     for f in ev.findings if f.kind != "uncached"],
    })
    return result
